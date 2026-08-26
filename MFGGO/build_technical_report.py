from __future__ import annotations

import os
import textwrap
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "MFGGO技术架构与代码深度分析报告.docx"
ASSET_DIR = ROOT / "tmp" / "technical-report-assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

FONT_PATH = ROOT / "public" / "fonts" / "SimHei.ttf"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.replace("#", "")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


NAVY = rgb("0B2545")
BLUE = rgb("2E74B5")
DARK_BLUE = rgb("1F4D78")
TEAL = rgb("0F6B78")
INK = rgb("18232F")
MUTED = rgb("5B6875")
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = rgb("9A6A00")
RISK = rgb("9B1C1C")
WHITE = "FFFFFF"


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Microsoft YaHei", size=11, color=INK, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.font.italic = italic
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def shade(element, fill: str):
    if hasattr(element, "_tc"):
        tc_pr = element._tc.get_or_add_tcPr()
    elif hasattr(element, "get_or_add_tcPr"):
        tc_pr = element.get_or_add_tcPr()
    else:
        tc_pr = element._p.get_or_add_pPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn("w:" + key), str(value))


def set_table_geometry(table, widths: Sequence[int], indent: int = 120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    old_grid = tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        old_grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[min(index, len(widths) - 1)]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_text(paragraph, text, *, bold=False, italic=False, color=INK, size=11, font="Microsoft YaHei"):
    run = paragraph.add_run(str(text))
    set_run_font(run, font, size, color, bold, italic)
    return run


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_body(doc, text, *, style="Normal", after=None, before=None, keep=False, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.add_run(text)
    if keep:
        set_keep_with_next(p)
    return p


def add_heading(doc, text, level=1, *, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    set_keep_with_next(p)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def add_bullets(doc, items: Iterable[str], *, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(style=style)
        p.add_run(item)


def add_numbered(doc, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_note(doc, label: str, text: str, *, color="0F6B78", fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=120)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_border(cell, left={"val": "single", "sz": 18, "color": color})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_text(p, label + "  ", bold=True, color=rgb(color), size=10.5)
    add_text(p, text, color=INK, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int], *, header_fill=LIGHT_BLUE, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths, indent=120)
    header = table.rows[0]
    repeat_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        shade(cell, header_fill)
        set_cell_border(cell, top={"val": "single", "sz": 6, "color": "B7C6D4"}, bottom={"val": "single", "sz": 6, "color": "B7C6D4"}, left={"val": "single", "sz": 4, "color": "D7E0E8"}, right={"val": "single", "sz": 4, "color": "D7E0E8"})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        clear_paragraph(p)
        add_text(p, value, bold=True, color=NAVY, size=font_size)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            if len(table.rows) % 2 == 0:
                shade(cell, "FBFCFD")
            set_cell_border(cell, top={"val": "single", "sz": 4, "color": "E2E8EE"}, bottom={"val": "single", "sz": 4, "color": "E2E8EE"}, left={"val": "single", "sz": 4, "color": "E2E8EE"}, right={"val": "single", "sz": 4, "color": "E2E8EE"})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            clear_paragraph(p)
            add_text(p, value, color=INK, size=font_size)
    # Re-apply the caller's geometry after all rows exist so every tcW agrees
    # with the intended column widths (including rows added after table creation).
    set_table_geometry(table, widths, indent=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, code: str, source: str, *, language="JavaScript"):
    p = doc.add_paragraph(style="Code Caption")
    add_text(p, f"代码摘录 | {language} | {source}", bold=True, color=TEAL, size=9)
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=120)
    cell = table.cell(0, 0)
    shade(cell, "F4F6F9")
    set_cell_border(cell, top={"val": "single", "sz": 4, "color": "D7E0E8"}, bottom={"val": "single", "sz": 4, "color": "D7E0E8"}, left={"val": "single", "sz": 4, "color": "D7E0E8"}, right={"val": "single", "sz": 4, "color": "D7E0E8"})
    lines = code.strip("\n").splitlines()
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.style = "Code Block"
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        clear_paragraph(p)
        add_text(p, line, font="Consolas", size=8.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(doc, path: Path, width=6.25, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    inline_shape = p.add_run().add_picture(str(path), width=Inches(width))
    # Set descriptive alt text for screen readers and accessibility audits.
    description = caption or path.stem
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)
    if caption:
        add_caption(doc, caption)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Microsoft YaHei")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=11, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    set_style_font(title, size=27, color=NAVY, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, size=14, color=MUTED)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True

    h1 = styles["Heading 1"]
    set_style_font(h1, size=16, color=BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True

    h2 = styles["Heading 2"]
    set_style_font(h2, size=13, color=BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    h3 = styles["Heading 3"]
    set_style_font(h3, size=12, color=DARK_BLUE, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        set_style_font(style, size=10.8, color=INK)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.10

    if "Code Block" not in [s.name for s in styles]:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    set_style_font(code_style, name="Consolas", size=8.2, color=INK)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0

    if "Code Caption" not in [s.name for s in styles]:
        code_caption = styles.add_style("Code Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_caption = styles["Code Caption"]
    set_style_font(code_caption, size=9, color=TEAL, bold=True)
    code_caption.paragraph_format.space_before = Pt(4)
    code_caption.paragraph_format.space_after = Pt(2)

    caption = styles["Caption"]
    set_style_font(caption, size=9, color=MUTED, italic=True)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    for style_name in ("TOC 1", "TOC 2", "TOC 3"):
        if style_name not in [s.name for s in styles]:
            st = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            st = styles[style_name]
        set_style_font(st, size=11 if style_name == "TOC 1" else 10.5, color=NAVY if style_name == "TOC 1" else MUTED)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.left_indent = Inches(0.25 * (int(style_name[-1]) - 1))

    header = section.header
    hp = header.paragraphs[0]
    clear_paragraph(hp)
    hp.paragraph_format.space_after = Pt(0)
    add_text(hp, "MFGGO 智造云  |  技术架构与代码深度分析", bold=True, color=MUTED, size=8.5)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = section.footer
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    add_text(fp, "内部技术资料  |  代码快照 2026-08-26  |  ", color=MUTED, size=8.5)
    add_page_number(fp)


def diagram_font(size: int):
    try:
        return ImageFont.truetype(str(FONT_PATH), size=size)
    except Exception:
        return ImageFont.load_default()


def draw_box(draw, xy, title, subtitle, fill, outline="#CBD7E2"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    title_font = diagram_font(30)
    sub_font = diagram_font(21)
    draw.text((x1 + 22, y1 + 18), title, fill="#102A43", font=title_font)
    lines = textwrap.wrap(subtitle, width=20)
    for idx, line in enumerate(lines):
        draw.text((x1 + 22, y1 + 60 + idx * 28), line, fill="#4B5D6B", font=sub_font)


def arrow(draw, start, end, color="#5D7285"):
    draw.line([start, end], fill=color, width=5)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 16
    left = (end[0] - size * math.cos(angle - 0.5), end[1] - size * math.sin(angle - 0.5))
    right = (end[0] - size * math.cos(angle + 0.5), end[1] - size * math.sin(angle + 0.5))
    draw.polygon([end, left, right], fill=color)


def make_architecture_diagram():
    image = Image.new("RGB", (1800, 980), "#F7F9FB")
    draw = ImageDraw.Draw(image)
    title_font = diagram_font(38)
    draw.text((70, 42), "MFGGO 运行时架构", fill="#0B2545", font=title_font)
    draw.text((72, 92), "云端协同能力与浏览器本地几何能力分层组合", fill="#5B6875", font=diagram_font(22))
    draw_box(draw, (80, 180, 520, 360), "多页面入口", "index / enterprise / platform / workspace / onlyoffice", "#E8F1F8")
    draw_box(draw, (650, 180, 1080, 360), "浏览器运行时", "zhizao-app.js · app.js · Three.js · PDF/OCR", "#EAF6F3")
    draw_box(draw, (1210, 180, 1710, 360), "外部编辑器", "ONLYOFFICE DocumentServer", "#FFF4D6")
    draw_box(draw, (260, 570, 760, 790), "Koa REST API", "认证 · 租户上下文 · authorize · 业务路由", "#E8EEF5")
    draw_box(draw, (980, 570, 1480, 790), "SQLite 数据层", "组织 · 项目 · 零件 · 报价 · FAIR · 审计", "#E8EEF5")
    draw_box(draw, (300, 850, 760, 940), "文件存储", "office-documents / chat-attachments", "#F2F4F7")
    draw_box(draw, (1020, 850, 1480, 940), "部署边界", "dist · systemd · 8320 / 4310", "#F2F4F7")
    arrow(draw, (520, 270), (650, 270))
    arrow(draw, (1080, 270), (1210, 270))
    arrow(draw, (850, 360), (620, 570))
    arrow(draw, (920, 360), (1110, 570))
    arrow(draw, (1210, 360), (1110, 570))
    arrow(draw, (500, 790), (500, 850))
    arrow(draw, (1230, 790), (1230, 850))
    return image.save(ASSET_DIR / "architecture_overview.png")


def make_data_flow_diagram():
    image = Image.new("RGB", (1800, 900), "#F7F9FB")
    draw = ImageDraw.Draw(image)
    draw.text((70, 42), "关键数据流", fill="#0B2545", font=diagram_font(38))
    draw.text((72, 92), "登录上下文、项目业务链与本地工程识别链彼此连接但不混写", fill="#5B6875", font=diagram_font(22))
    draw_box(draw, (80, 190, 480, 340), "1. 登录", "Bearer token -> sessions -> organization_id", "#E8F1F8")
    draw_box(draw, (660, 190, 1060, 340), "2. 项目", "projects -> project_members -> parts", "#EAF6F3")
    draw_box(draw, (1240, 190, 1720, 340), "3. 业务结果", "quote / FAIR / task / message / audit", "#FFF4D6")
    arrow(draw, (480, 265), (660, 265))
    arrow(draw, (1060, 265), (1240, 265))
    draw_box(draw, (80, 540, 480, 690), "CAD 输入", "STEP / IGES / BREP -> OpenCascade", "#E8F1F8")
    draw_box(draw, (660, 540, 1060, 690), "几何特征", "mesh metrics -> holes / curved faces / thin wall", "#EAF6F3")
    draw_box(draw, (1240, 540, 1720, 690), "报价导出", "ExcelJS 5-sheet workbook + jsPDF", "#FFF4D6")
    arrow(draw, (480, 615), (660, 615))
    arrow(draw, (1060, 615), (1240, 615))
    draw.text((80, 770), "PDF 质检分支: pdf.js vector text -> visual symbol analysis -> Tesseract -> engineering parser -> FAIR rows", fill="#36566F", font=diagram_font(24))
    return image.save(ASSET_DIR / "data_flow.png")


def source_line_count(name: str) -> int:
    try:
        return len((ROOT / name).read_text(encoding="utf-8").splitlines())
    except Exception:
        return 0


def build_report():
    make_architecture_diagram()
    make_data_flow_diagram()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "MFGGO 智造云技术架构与代码深度分析报告"
    doc.core_properties.subject = "机加工 3D 零件智能报价助手"
    doc.core_properties.author = "技术分析"
    doc.core_properties.comments = "基于 2026-08-26 工作区源码快照生成"

    # Cover page
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.paragraph_format.space_after = Pt(10)
    add_text(p, "TECHNICAL ARCHITECTURE REPORT", bold=True, color=TEAL, size=10.5)
    p = doc.add_paragraph(style="Title")
    add_text(p, "MFGGO 智造云", bold=True, color=NAVY, size=29)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(12)
    add_text(p, "机加工 3D 零件智能报价助手", bold=True, color=NAVY, size=27)
    p = doc.add_paragraph(style="Subtitle")
    add_text(p, "技术架构与代码深度分析报告", color=MUTED, size=15)
    add_note(doc, "核心判断", "当前系统不是单纯的页面原型，而是由浏览器本地 CAD/FAIR 引擎与云端多租户协同工作台组成的双运行时产品。云端 P0 权限边界已经具备可验收基础，但距离公网生产级 SaaS 仍缺少成员生命周期、可配置 RBAC、强认证、对象存储、可恢复备份和报价版本审批。", color="0F6B78")
    meta_rows = [
        ("代码快照", "2026-08-26；工作区无 Git 元数据，以当前文件为准"),
        ("版本基线", "package.json 1.1.0；Node.js ESM；Vite 多页面"),
        ("分析范围", "前端入口、云端工作台、Koa API、SQLite schema、CAD/FAIR/OCR、ONLYOFFICE、测试与部署"),
        ("验证方式", "源码静态分析 + Node 内置测试；82 个测试用例全部通过"),
        ("读者对象", "架构师、后端/前端工程师、制造数字化负责人、部署与验收人员"),
    ]
    add_table(doc, ["项目", "说明"], meta_rows, [1700, 7660], font_size=9.6)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, "内部技术资料  ·  供架构评审、开发交接与上线前整改使用", color=MUTED, size=9.5, italic=True)
    doc.add_page_break()

    # Manual contents
    add_heading(doc, "目录", 1)
    add_body(doc, "本目录对应 Word Heading 1/2/3 样式。打开 Word 导航窗格即可按标题跳转；代码摘录后的路径和行号用于回到源码定位。", after=10)
    toc = [
        ("1  文档说明与结论", 1), ("2  系统定位与边界", 1), ("3  总体架构", 1),
        ("4  运行与部署", 1), ("5  云端多租户业务链", 1), ("6  数据模型与持久化", 1),
        ("7  API 与后端代码分析", 1), ("8  前端云工作台分析", 1),
        ("9  浏览器本地 CAD/FAIR/报价链", 1), ("10  ONLYOFFICE 文件链", 1),
        ("11  安全、权限与审计", 1), ("12  测试与质量", 1),
        ("13  风险、技术债与改进路线", 1), ("14  附录", 1),
    ]
    for title, level in toc:
        p = doc.add_paragraph(style="TOC 1")
        add_text(p, title, color=NAVY, size=11)
    add_note(doc, "阅读建议", "先读第 3、5、7、9 节建立系统心智模型，再读第 11、13 节判断上线风险。第 14 节给出逐文件索引、接口矩阵和本地验收命令。", color="7A5A00", fill="FFF8E8")
    doc.add_page_break()

    # 1
    add_heading(doc, "1  文档说明与结论", 1)
    add_heading(doc, "1.1 分析方法与证据等级", 2)
    add_body(doc, "本报告以工作区 2026-08-26 的源码快照为事实来源。由于目录没有 Git 元数据，文档不把提交历史、分支状态或未在文件中出现的外部服务视为已验证事实。结论分为三类：已实现表示可在代码或测试中直接定位；已声明表示 README/ARCHITECTURE.md 有明确说明但仍需部署条件；建议项表示为正式生产或规模化运行提出的改进，不代表当前已经存在。")
    add_note(doc, "代码摘录说明", "代码块只摘取与分析结论直接相关的路径；其中的 `...` 表示省略非关键参数或上下文，括号内文件与行号仍可回到完整源码核对。", color="0F6B78", fill="EAF6F3")
    evidence_rows = [
        ("已实现", "Koa REST API、SQLite 持久化、组织隔离、角色/模块/项目成员授权、项目/零件/报价/FAIR/任务/沟通/文档/审计", "server/app.js; server/db.js; tests/saas-api.test.mjs"),
        ("已实现", "浏览器端 STEP/IGES/BREP 解析、Three.js 预览、PDF 框选、视觉符号识别、OCR、Excel/PDF 导出", "app.js; inspection-recognition.js; inspection-pipeline.js"),
        ("依赖条件", "ONLYOFFICE 在线编辑、公网 8320 systemd 部署、独立 DocumentServer 8081", "README.md; deploy/*; server/app.js"),
        ("尚未落地", "订单、客户/供应商、工艺路线、审批流、排产、交付、报价版本、订阅计费", "ARCHITECTURE.md:241-259"),
    ]
    add_table(doc, ["证据等级", "结论", "主要证据"], evidence_rows, [1200, 5200, 2960], font_size=9)
    add_heading(doc, "1.2 一句话架构结论", 2)
    add_body(doc, "系统采用“多入口前端 + 双运行时 + 单体 API + SQLite/文件目录”的组合：Vite 构建多个 HTML 入口；云端工作台通过 zhizao-app.js 调用 Koa REST API；独立 workspace.html 在浏览器中加载 Three.js、OpenCascade WASM、pdf.js、Tesseract.js 和 ExcelJS；Koa 侧把业务数据写入 SQLite，把 Office 文件和聊天附件写入按组织划分的目录，并通过 ONLYOFFICE 提供在线编辑。")
    add_note(doc, "架构评价", "对于本地部署、内部验收和小规模单实例场景，这一组合的优点是安装简单、图纸默认不出浏览器、业务接口边界清晰。对于公网多实例和正式商业化场景，SQLite、文件目录、静态角色表和缺少强认证会成为主要约束。", color="9A6A00", fill="FFF8E8")
    add_heading(doc, "1.3 当前能力与目标能力的分界", 2)
    add_body(doc, "代码已经把“可运行的多租户协同基础”做出来了，但业务链仍停在项目、报价和质量准备阶段。ARCHITECTURE.md 中描述的完整链路是“客户询盘 -> 项目 -> 图纸/3D -> 工艺材料/BOM -> 成本核算与报价版本 -> 审批/发送 -> 订单排产 -> 加工/FAIR -> 交付”。当前持久化模型覆盖到报价草稿、FAIR 检验项和沟通审计，尚未覆盖订单及其后续执行对象。")

    # 2
    add_heading(doc, "2  系统定位与边界", 1, page_break=True)
    add_heading(doc, "2.1 产品分成两个互补工作台", 2)
    add_table(doc, ["工作台", "入口", "主要职责", "数据边界"], [
        ("企业云工作台", "index.html / enterprise.html", "项目看板、任务、零件中心、制表中心、BOM 报价、项目沟通、企业聊天、统计", "数据写入 SQLite；按 organization_id 和项目成员隔离"),
        ("平台主后台", "admin.html / platform.html", "企业管理、平台用户、模块开关、跨企业审计、企业上下文切换", "仅 platform_admin；切换后创建带 organization_id 的新会话"),
        ("3D/FAIR 工作台", "workspace.html", "上传 CAD/PDF、几何解析、工程图识别、FAIR 表、报价基础和导出", "默认只在浏览器内存中处理；不自动写云端业务表"),
        ("在线文档编辑", "onlyoffice.html", "在文档库中打开 Word/Excel/PPT/PDF 并回写文件", "文件由 Koa 签名下载和回调保存；编辑能力由权限决定"),
    ], [1700, 1800, 3500, 2360], font_size=8.9)
    add_heading(doc, "2.2 业务边界", 2)
    add_body(doc, "云端模型把“项目”作为大多数业务资源的聚合边界，零件、报价、FAIR、任务、项目会话和 Office 文档都可以关联 project_id。平台模型把“企业”作为租户边界，登录会话同时绑定 user_id 和 organization_id。浏览器本地模型则以本次上传的文件集合为边界，state.files 保存 CAD/PDF、解析结果、材料、数量、报价和图纸关联关系。")
    add_bullets(doc, [
        "隐私边界：STEP/PDF 默认不上传服务器，适合客户图纸敏感的内网或本地使用；云端文档和聊天附件则会写入服务器目录。",
        "业务边界：报价算法是几何和成本的估算基线，不是最终工艺定额；FAIR 识别结果需要人工复核，代码保留 source、ocrRaw、confidence 和 manual 标记。",
        "平台边界：套餐、计费、价格策略、客户/供应商主数据、订单和排产不在当前实现范围。",
    ])
    add_heading(doc, "2.3 关键架构原则", 2)
    add_table(doc, ["原则", "代码体现", "工程含义"], [
        ("服务端是权限权威", "authorize(ctx, permission, { moduleKey, projectId })", "前端隐藏导航只改善体验，不能替代角色、模块和项目范围校验。"),
        ("组织隔离优先", "业务表普遍带 organization_id；查询先取 session.organizationId", "客户端不能通过传入租户 ID 读取其他企业。"),
        ("本地计算优先", "workspace.html 直接加载 occt-import-js.wasm、pdf.js、Tesseract", "客户图纸可在浏览器完成解析，降低上传和泄露风险。"),
        ("证据可追溯", "audit_events、quote lines、FAIR source/ocrEvidence", "修改、报价和识别结果都能回到操作者或识别来源。"),
    ], [1900, 3500, 3960], font_size=9)

    # 3
    add_heading(doc, "3  总体架构", 1, page_break=True)
    add_image(doc, ASSET_DIR / "architecture_overview.png", width=6.35, caption="图 1  运行时架构：云端协同与浏览器本地工程能力分层")
    add_heading(doc, "3.1 分层说明", 2)
    add_table(doc, ["层", "实现", "关键职责", "主要风险"], [
        ("入口层", "6 个 Vite HTML 入口", "登录、企业/平台导航、3D/FAIR、ONLYOFFICE 容器", "多入口共享状态有限，跨页面切换依赖 sessionStorage/postMessage。"),
        ("云端客户端层", "zhizao-app.js + cloud.css", "store、hydrate、渲染、事件委托、API 调用、模块化导航", "大量 innerHTML 和单文件逻辑使长期维护成本上升。"),
        ("本地工程层", "app.js + inspection-*", "CAD 解析、网格指标、特征推断、PDF/OCR、FAIR 和报价导出", "受浏览器内存/CPU、WASM 和 OCR 质量影响。"),
        ("API 层", "Koa 3 + @koa/router", "认证、输入校验、权限、业务路由、文件回调、统一错误", "路由和 DB 调用集中在两个大文件，事务边界不够明显。"),
        ("数据层", "node:sqlite DatabaseSync", "schema、迁移、seed、查询映射、审计", "单实例写入和备份恢复能力有限。"),
        ("文件/集成层", "本地目录 + ONLYOFFICE", "Office 文件、聊天附件、签名 URL、保存回调", "对象存储、病毒扫描、签名过期和多实例共享尚未完成。"),
    ], [1300, 2200, 3600, 2260], font_size=8.7)
    add_heading(doc, "3.2 典型运行时路径", 2)
    add_numbered(doc, [
        "用户在任一云端入口提交登录表单，POST /api/auth/login 校验 scrypt 密码并创建 12 小时 session token。",
        "前端把 token 放入 sessionStorage，GET /api/me 恢复用户、企业、角色、平台管理员标志和模块开关。",
        "hydrateProjects 并行加载项目、零件、成员、任务、会话、统计和文档，再按当前首个项目补充报价和 FAIR。",
        "用户写入业务对象时，Koa 路由先做角色权限、模块开关、企业边界和 project_members 校验，再调用 db service 写 SQLite，并追加 audit_events。",
        "用户进入 3D/FAIR 工作台时，文件留在浏览器内存；CAD 走 OpenCascade + Three.js，PDF 走 vector text + visual symbol + OCR + engineering parser。",
    ])
    add_heading(doc, "3.3 架构取舍", 2)
    add_body(doc, "当前实现选择了“少服务、少部署、强客户端”的路线。它把几何解析和复杂 PDF 识别放在浏览器，避免后端引入 CAD 运行时和大文件队列；把协同、权限、审计和文档放在 Koa/SQLite，便于验收和演示。代价是浏览器设备差异会影响解析体验，云端报价与本地报价存在两套模型，后端单体文件在扩展到更多制造域时会迅速变大。")

    # 4
    add_heading(doc, "4  运行与部署", 1, page_break=True)
    add_heading(doc, "4.1 技术栈与脚本", 2)
    add_table(doc, ["类别", "当前实现", "证据"], [
        ("语言/模块", "Node.js ESM；浏览器端原生 ES modules", "package.json:1-5"),
        ("前端构建", "Vite 7.1.1；base='./'；6 个 Rollup input；/api 代理", "vite.config.js:4-21"),
        ("服务端", "Koa 3、@koa/router、koa-bodyparser、koa-static", "package.json:14-24; server/app.js:1-8"),
        ("本地几何", "Three.js 0.185、occt-import-js 0.0.23、Tesseract.js 7", "package.json:21-28; workspace.html"),
        ("导出", "ExcelJS、jsPDF、docx、pptxgenjs", "package.json:17-29"),
        ("数据库", "node:sqlite DatabaseSync + 外键 + 增量 ALTER + 索引", "server/db.js:1-7, 354-381"),
    ], [1800, 5100, 2460], font_size=9)
    add_code(doc, """const root = process.cwd();
const host = process.env.HOST || '127.0.0.1';
const port = Number(process.env.PORT || 4310);
const staticDir = join(root, 'dist');
const dbPath = process.env.DB_PATH || join(root, 'data', 'machquote.sqlite');

const app = await createApp({ dbPath, staticDir });
const server = app.listen(port, host, () => {
  console.log(`MFGGO SaaS API listening on http://${host}:${port}`);
});""", "server/index.js:5-27", language="JavaScript")
    add_body(doc, "server/index.js 在非回环地址启动时要求 PUBLIC_BASE_URL、ONLYOFFICE_JWT_SECRET，并在首次初始化时要求 INITIAL_ADMIN_PASSWORD。这是一个有价值的启动闸门，但 app.js 内部仍保留了本地默认 JWT secret，不能把启动闸门当作完整的密钥治理。", after=6)
    add_heading(doc, "4.2 开发、预览与公网模式", 2)
    add_table(doc, ["模式", "命令/入口", "端口与数据", "适用场景"], [
        ("开发", "pnpm run dev", "API 4310 + Vite 4173；Vite 代理 /api", "联调、快速验收"),
        ("分别调试", "pnpm run dev:server / pnpm run dev:web", "服务和前端独立启动", "排查 API 或 UI 问题"),
        ("本地预览", "pnpm run build; pnpm run server", "静态 dist 由 Koa 提供；默认 4310", "接近生产的单机验收"),
        ("独立公网", "systemd zhizao-cloud", "0.0.0.0:8320；DB_PATH=/var/lib/zhizao-cloud", "内网/受控公网；仍需 HTTPS/VPN"),
    ], [1500, 2600, 3000, 2260], font_size=9)
    add_heading(doc, "4.3 配置项与运营意义", 2)
    config_rows = [
        ("HOST / PORT", "监听地址与端口；公网示例为 0.0.0.0 / 8320", "公开监听必须配合防火墙和 HTTPS"),
        ("DB_PATH", "SQLite 文件路径，默认 data/machquote.sqlite", "需要备份、锁文件和迁移策略"),
        ("INITIAL_ADMIN_PASSWORD", "首次 seed 的管理员密码；README 默认 123456", "必须在公网初始化后删除/替换"),
        ("PUBLIC_BASE_URL", "ONLYOFFICE 可访问的 API 根地址", "必须使用外部可达、协议正确的 URL"),
        ("ONLYOFFICE_JWT_SECRET", "服务端 JWT/HMAC 密钥", "必须用长随机值，禁止复用示例值"),
        ("ONLYOFFICE_STORAGE_DIR / CHAT_ATTACHMENT_STORAGE_DIR", "按组织保存 Office 文件和聊天附件", "生产应迁移对象存储并做病毒扫描"),
    ]
    add_table(doc, ["变量", "用途", "上线注意"], config_rows, [2300, 3900, 3160], font_size=8.9)
    add_heading(doc, "4.4 systemd 部署分析", 2)
    add_body(doc, "deploy/zhizao-cloud.service 使用独立用户、NoNewPrivileges、PrivateTmp、ProtectHome、ProtectSystem=strict，并仅通过 ReadWritePaths 放行 /var/lib/zhizao-cloud。这些是合理的最小权限基线。README 仍建议直接 IP 明文 HTTP 访问，这只能用于内部验收，不应承载正式客户图纸、报价或凭证。上线前应把 8320 放在 HTTPS 反向代理或 VPN 后面，限制管理入口和 DocumentServer 网络范围。")
    add_code(doc, """[Service]
User=zhizao-cloud
WorkingDirectory=/opt/zhizao-cloud
EnvironmentFile=/etc/zhizao-cloud/zhizao-cloud.env
ExecStart=/usr/bin/node /opt/zhizao-cloud/server/index.js
Restart=on-failure
NoNewPrivileges=true
ProtectHome=true
ProtectSystem=strict
ReadWritePaths=/var/lib/zhizao-cloud""", "deploy/zhizao-cloud.service:6-19", language="systemd")

    # 5
    add_heading(doc, "5  云端多租户业务链", 1, page_break=True)
    add_heading(doc, "5.1 导航与模块开关", 2)
    add_table(doc, ["模块键", "界面标签", "典型 API", "当前权限要点"], [
        ("projects", "项目管理", "/api/projects", "项目读写；非管理员依赖 project_members"),
        ("tasks", "我的任务", "/api/tasks", "项目任务按项目范围；独立任务默认只能操作本人"),
        ("parts", "零件中心", "/api/parts", "零件读写；FAIR 也复用 parts 模块开关"),
        ("workspace", "制表中心", "/api/documents", "Office 文档权限、项目范围、回收站"),
        ("bom", "BOM 报价助手", "/api/projects/:id/quotes", "工程师可写报价，QA 默认无 quote 权限"),
        ("communication", "项目沟通", "/api/conversations?projectId=...", "项目会话需要项目成员关系"),
        ("chat", "聊天信息", "/api/conversations", "企业会话和附件按 chat 权限"),
        ("stats", "统计", "/api/stats", "当前组织聚合统计"),
    ], [1500, 1900, 2900, 3060], font_size=8.8)
    add_heading(doc, "5.2 项目看板与资源聚合", 2)
    add_body(doc, "zhizao-app.js 将项目阶段规范化为十个看板列，例如“立项沟通、询盘发布、内部报价、对外报价、订单发布/待办、已排产/处理中、异常/优先处理、已到货/质检、发货、订单结束/已完成”。这套映射主要是前端展示语义；服务端 projects.stage 仍是可变字符串，因此后续若要驱动审批、排产或统计，必须把阶段变成受约束的状态字典而不是继续依赖正则归类。")
    add_body(doc, "GET /api/projects/:projectId/workspace 是云端项目卡片的聚合接口。它先用 workspace 模块和 project.read 校验项目，再按 modules 与 roleCan 为 tasks、parts、quotes、fairItems、documents、conversations、comments、activities 分别决定是否返回。这个设计避免了前端并行请求时把不该看到的报价或审计数据混入统一响应。")
    add_code(doc, """const modules = db.listOrganizationModules(session.organizationId);
const canRead = (permission, moduleKey) =>
  modules[moduleKey] && roleCan(session.role, permission);

ctx.body = {
  project,
  tasks: canRead('task.read', 'tasks') ? db.listProjectTasks(...) : [],
  parts: canRead('part.read', 'parts') ? db.listParts(...) : [],
  quotes: canRead('quote.read', 'bom') ? db.listQuotes(...) : [],
  fairItems: canRead('fair.read', 'parts') ? db.listFairItems(...) : [],
  documents: canRead('document.read', 'workspace') ? db.listOfficeDocuments(...) : [],
  conversations: canRead('communication.read', 'communication') ? db.listProjectConversations(...) : [],
  activities: roleCan(session.role, 'audit.read') ? db.listProjectActivity(...) : []
};""", "server/app.js:589-605", language="JavaScript")
    add_heading(doc, "5.3 前端云端数据流", 2)
    add_code(doc, """const [projectsResult, partsResult, membersResult,
  tasksResult, conversationsResult, statsResult] = await Promise.all([
  apiRequest('/api/projects', {}, store.apiToken),
  apiRequest('/api/parts', {}, store.apiToken),
  apiRequest('/api/members', {}, store.apiToken),
  apiRequest('/api/tasks', {}, store.apiToken),
  apiRequest('/api/conversations', {}, store.apiToken),
  apiRequest('/api/stats', {}, store.apiToken)
]);
...
const projectId = store.projects[0]?.id;
if (projectId) {
  store.quotes = (await apiRequest(`/api/projects/${projectId}/quotes`, ...)).body.quotes;
  store.fairItems = (await apiRequest(`/api/projects/${projectId}/fair-items`, ...)).body.items;
}""", "zhizao-app.js:758-797", language="JavaScript")
    add_body(doc, "这里有一个可预见的状态演进问题：首次 hydrate 只主动补充第一个项目的报价和 FAIR；其他项目依赖后续弹窗或工作台请求。如果把项目列表、报价和 FAIR 做成可切换的深层页面，建议统一改成按 projectId 缓存的 resource store，避免在项目切换时复用旧报价。")

    # 6
    add_heading(doc, "6  数据模型与持久化", 1, page_break=True)
    add_image(doc, ASSET_DIR / "data_flow.png", width=6.35, caption="图 2  登录、项目业务链与本地工程识别链的数据流")
    add_heading(doc, "6.1 表分组与关系", 2)
    schema_rows = [
        ("租户与身份", "organizations, users, memberships, organization_modules, platform_admins, sessions", "用户、企业、成员角色、模块开关、平台管理员、会话"),
        ("项目与制造资料", "projects, project_members, parts, files", "项目阶段、项目成员、零件元数据、通用文件"),
        ("Office 文档", "office_documents, office_document_permissions", "文档库、项目关联、权限、版本、回收站"),
        ("报价与质量", "quotes, quote_lines, fair_items", "报价头/行、分为单位的价格、FAIR 特性/名义值/公差/状态"),
        ("任务与沟通", "tasks, task_subtasks, task_comments, project_comments", "任务、子任务、任务评论、项目评论"),
        ("聊天与附件", "conversations, messages, conversation_user_states, chat_attachments", "企业/项目会话、消息、已读/稍后、附件"),
        ("审计", "audit_events", "企业范围和平台范围的操作事件"),
    ]
    add_table(doc, ["域", "表", "用途"], schema_rows, [1700, 4400, 3260], font_size=8.8)
    add_heading(doc, "6.2 租户隔离的实现方式", 2)
    add_body(doc, "schema 中除 users、platform_admins 等身份表外，业务表均包含 organization_id；项目资源进一步包含 project_id。db service 的查询通常把 organization_id 放在 WHERE 条件的第一组约束中，路由层从 ctx.state.session 读取当前组织，不信任请求体中的 organizationId。外键启用 ON DELETE CASCADE 或 SET NULL，createDatabase 还补充了 project、task、message、attachment 等高频索引。")
    add_table(doc, ["不变量", "当前实现", "审查结论"], [
        ("组织归属", "业务查询按 organization_id 过滤", "基础隔离正确；应增加统一 DAO/查询断言，避免新表漏列"),
        ("项目访问", "owner/admin 全项目；其他角色必须在 project_members", "P0 已覆盖主要 API"),
        ("资源关联", "quotes/fair_items 的 project_id 通过 authorize 校验", "项目边界正确；partId 关联仍需做同组织/同项目校验"),
        ("文件路径", "存储 key 由组织 ID + 时间戳/随机串生成，文件名只用于元数据", "路径穿越风险有所缓解；仍需内容扫描和对象存储"),
        ("审计", "写操作调用 addAudit，列表最多返回 100/200 条", "可追溯基线完成；长期需分页、不可篡改和归档"),
    ], [1800, 3900, 3660], font_size=8.8)
    add_code(doc, """CREATE TABLE IF NOT EXISTS projects (
  id TEXT PRIMARY KEY,
  organization_id TEXT NOT NULL REFERENCES organizations(id) ON DELETE CASCADE,
  title TEXT NOT NULL,
  stage TEXT NOT NULL,
  owner_user_id TEXT REFERENCES users(id),
  progress INTEGER NOT NULL DEFAULT 0 CHECK(progress BETWEEN 0 AND 100),
  created_by TEXT REFERENCES users(id),
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS project_members (
  organization_id TEXT NOT NULL REFERENCES organizations(id) ON DELETE CASCADE,
  project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
  user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  project_role TEXT NOT NULL DEFAULT 'member',
  PRIMARY KEY(project_id, user_id)
);""", "server/db.js:41-68", language="SQL")
    add_heading(doc, "6.3 密码、会话与迁移", 2)
    add_code(doc, """export function hashPassword(password, salt = randomBytes(16).toString('hex')) {
  return `${salt}:${scryptSync(password, salt, 32).toString('hex')}`;
}

export function verifyPassword(password, encoded) {
  const [salt, expected] = String(encoded || '').split(':');
  const actual = scryptSync(password, salt, 32);
  return expectedBuffer.length === actual.length
    && timingSafeEqual(actual, expectedBuffer);
}""", "server/db.js:252-262", language="JavaScript")
    add_body(doc, "会话 token 使用 randomBytes(32) 生成，过期时间固定为 12 小时；注销会删除当前 session，但没有设备列表、全局撤销、刷新令牌、异常登录检测或定时清理。createDatabase 采用 try/catch 包裹 ALTER TABLE 的增量迁移，适合演示和小范围迭代，但生产应采用显式 migration table、版本号和失败回滚。")

    # 7
    add_heading(doc, "7  API 与后端代码分析", 1, page_break=True)
    add_heading(doc, "7.1 Koa 组装与错误边界", 2)
    add_body(doc, "server/app.js 的 createApp 同时负责数据库初始化、Koa 中间件、路由、文件根路径和 ONLYOFFICE 配置。中间件顺序是异常捕获 -> bodyParser -> router -> allowedMethods -> koa-static。异常处理把非 expose 错误统一映射为“服务暂时不可用”，避免把内部堆栈直接返回前端；路由内部则使用 jsonError 返回结构化 error code/message。")
    add_code(doc, """function authRequired(db) {
  return async (ctx, next) => {
    const header = ctx.get('authorization');
    const token = header.startsWith('Bearer ')
      ? header.slice(7).trim() : '';
    const session = token ? db.getSession(token) : null;
    if (!session) return jsonError(ctx, 401,
      'unauthorized', '登录已失效，请重新登录');
    ctx.state.session = session;
    await next();
  };
}""", "server/app.js:10-20", language="JavaScript")
    add_heading(doc, "7.2 统一授权策略", 2)
    add_code(doc, """const authorize = (ctx, permission, { moduleKey = '', projectId = '' } = {}) => {
  const session = ctx.state.session;
  if (!roleCan(session.role, permission)) {
    return jsonError(ctx, 403, 'permission_denied', ...);
  }
  if (moduleKey && !db.listOrganizationModules(session.organizationId)[moduleKey]) {
    return jsonError(ctx, 403, 'module_disabled', ...);
  }
  if (!projectId) return true;
  const project = db.getProject(session.organizationId, projectId);
  if (!project) return jsonError(ctx, 404, 'project_not_found', ...);
  if (!hasProjectAccess(session, project.id)) {
    return jsonError(ctx, 403, 'project_access_denied', ...);
  }
  return project;
};""", "server/app.js:86-109", language="JavaScript")
    add_body(doc, "该策略把角色、模块和项目范围集中到一个调用点，拒绝码也有区分度，测试可以分别验证 permission_denied、module_disabled 和 project_access_denied。其局限是 ROLE_PERMISSIONS 是代码常量，权限命名粒度仍偏粗，无法表达报价成本字段、部门范围、审批动作或外部协作者。")
    add_heading(doc, "7.3 API 分组矩阵", 2)
    api_rows = [
        ("认证", "POST /auth/login; POST /auth/logout; GET /me", "session", "登录、注销、恢复组织上下文"),
        ("平台", "GET/POST /platform/organizations; PUT .../modules; POST .../switch", "platform_admin", "企业、模块和上下文切换"),
        ("项目", "GET/POST /projects; PUT /projects/:id; GET /projects/:id/workspace", "project.* + projects", "项目、聚合工作台"),
        ("项目成员", "GET/POST /projects/:id/members", "project.member.manage", "项目范围授权"),
        ("零件", "GET /parts; GET/POST /projects/:id/parts; PUT /parts/:id", "part.* + parts", "零件元数据"),
        ("报价", "GET/POST/PUT /projects/:id/quotes...", "quote.* + bom", "报价头、BOM 行、分为单位金额"),
        ("FAIR", "GET/POST/PUT /projects/:id/fair-items...", "fair.* + parts", "检验特性与状态"),
        ("任务", "GET/POST/PUT /tasks...; subtasks/comments", "task.* + tasks", "任务、子任务、评论"),
        ("沟通", "GET/POST /conversations.../messages; PUT /state", "communication/chat", "项目会话、企业聊天、已读状态"),
        ("文件", "GET/POST /documents; upload; download; config", "document.* + workspace", "Office 文件库、回收站、编辑配置"),
        ("审计", "GET /audit; GET /platform/audit", "audit.read / platform_admin", "企业审计和平台审计"),
    ]
    add_table(doc, ["域", "接口", "授权", "功能"], api_rows, [1300, 3700, 1900, 2460], font_size=8.5)
    add_heading(doc, "7.4 代码级优点与缺口", 2)
    add_table(doc, ["观察", "证据", "影响"], [
        ("输入长度有统一上限", "validString；JSON 2 MB；文件 50 MB；消息 4000 字符", "降低异常输入和资源耗尽风险"),
        ("文件上传走原始流", "documents/upload、chat-attachments 逐 chunk 累计", "避免把大文件交给 JSON parser；仍需超时/并发限流"),
        ("报价写入缺少显式事务", "createQuote/updateQuote 先写头再循环写行", "中途异常可能留下不完整报价；应使用 BEGIN/COMMIT/ROLLBACK"),
        ("关联 ID 校验不完整", "quote_lines.part_id、fair_items.part_id 接受请求值", "需验证 part 属于当前 organization 且属于 project，避免跨域关联"),
        ("审计列表固定上限", "listAudit LIMIT 100；平台 LIMIT 200", "长期审计会丢失旧记录的可查询性；需分页/归档"),
        ("大文件回调缺少细粒度保护", "ONLYOFFICE callback 直接 fetch body.url 并写回", "需限制来源、大小、超时、重试和幂等"),
    ], [2200, 3600, 3560], font_size=8.8)

    # 8
    add_heading(doc, "8  前端云工作台分析", 1, page_break=True)
    add_heading(doc, "8.1 store 与页面模式", 2)
    add_code(doc, """export function createCloudStore() {
  return {
    authenticated: false,
    platformAdmin: false,
    consoleMode: 'enterprise',
    currentView: 'projects',
    projects: [], parts: [], tasks: [],
    conversations: [], messages: {},
    quotes: [], fairItems: [],
    apiToken: '', user: null, organization: null,
    nav: NAV_ITEMS.map(item => ({ ...item }))
  };
}""", "zhizao-app.js:66-68", language="JavaScript")
    add_body(doc, "企业页和平台页复用同一个 zhizao-app.js，通过 body[data-console-mode] 在初始化时选择 NAV_ITEMS 或 PLATFORM_NAV_ITEMS。页面状态是单一内存 store，DOM 由 renderShell 和多个 render* 函数重建，交互使用 content.addEventListener 的事件委托。这个模式适合快速迭代和多入口复用，但状态与视图耦合在一个大模块中，后续应按领域拆分 reducer/resource cache 和 view renderer。")
    add_heading(doc, "8.2 登录恢复与失败处理", 2)
    add_body(doc, "restoreSession 从 sessionStorage 读取 mfggo_saas_token，先调用 /api/me，再调用 hydrateProjects。若账号不是平台管理员却打开 platform.html，前端清空状态并拒绝进入 shell；如果登录成功但企业数据加载失败，也会 resetClientState 并把错误显示在登录表单。这避免了“半登录”状态把旧租户数据留在页面上，是当前前端实现中较重要的安全细节。")
    add_code(doc, """const hydrated = await hydrateProjects(store);
if (!hydrated) {
  resetClientState();
  saveSession(store);
  error.textContent = '登录成功，但企业数据暂时无法加载，请检查服务后重试';
  error.classList.add('show');
  return;
}
showShell();""", "zhizao-app.js:1218-1234", language="JavaScript")
    add_heading(doc, "8.3 渲染、输入与异步竞态", 2)
    add_bullets(doc, [
        "esc() 对 innerHTML 插值做了 &、<、>、引号和单引号转义，降低项目名、消息、文件名造成的 XSS 风险。",
        "apiRequest() 统一注入 content-type 和 Bearer token，并将网络异常转换为 network_error，调用方再决定 toast 或回退。",
        "聊天和项目沟通加载保存 requestConversationId，并比较当前选中会话，避免旧请求返回后覆盖新会话。",
        "退出或新账号登录会清空项目、会话、附件草稿、文档列表和弹窗 DOM，测试对这些清理动作有源码断言。",
    ])
    add_heading(doc, "8.4 前端维护性判断", 2)
    add_body(doc, f"当前 zhizao-app.js 约 {source_line_count('zhizao-app.js')} 行，app.js 约 {source_line_count('app.js')} 行；两者都把大量渲染、事件、数据转换和业务规则放在单文件中。app.js 还包含多次 const renderX = renderX 包装式扩展，这能快速叠加功能，但会增加调用链追踪难度。建议下一阶段按 cloud/auth、cloud/projects、cloud/tasks、cloud/chat、cloud/documents、local/cad、local/inspection、local/quote 拆分，并为每个模块保留纯函数测试。")

    # 9
    add_heading(doc, "9  浏览器本地 CAD/FAIR/报价链", 1, page_break=True)
    add_heading(doc, "9.1 文件进入与安全阈值", 2)
    add_table(doc, ["阈值/检查", "实现", "目的"], [
        ("单文件 50 MB", "maxUploadBytes = 50 * 1024 * 1024", "避免浏览器一次加载超大模型或图纸"),
        ("ZIP 最多 120 项", "maxArchiveEntries = 120", "限制压缩包目录规模"),
        ("ZIP 解压 150 MB", "maxArchiveBytes = 150 * 1024 * 1024", "降低压缩炸弹风险"),
        ("路径穿越检查", "忽略 ..、目录项和 __MACOSX", "避免把归档路径当作真实文件路径"),
        ("密码 ZIP 拒绝", "flags & 1 -> error", "避免无法验证的加密条目"),
        ("支持格式", "STEP/STP/IGES/IGS/BREP/PDF/ZIP", "明确工程输入边界"),
    ], [2100, 4000, 3260], font_size=9)
    add_heading(doc, "9.2 CAD 解析与 B-Rep 特征推断", 2)
    add_body(doc, "parseCadItem 通过 ensureCadKernel 懒加载 OpenCascade WASM，按扩展名调用 ReadStepFile、ReadIgesFile 或 ReadBrepFile，单位设为 millimeter，线性偏差使用 bounding_box_ratio 0.001。cadMetrics 遍历 mesh 的 position/index，计算包围盒尺寸、三角形数、B-Rep 面数、表面积和有符号体积的绝对值，再交给 analyzeBrepFeatures。")
    add_table(doc, ["指标", "计算/推断", "用于"], [
        ("包围盒体积", "size.x * size.y * size.z 与实体体积取最大", "材料利用率和毛坯重量基线"),
        ("材料利用率", "实体体积 / 包围盒体积，限制在 0.01-1", "去除率与机时估算"),
        ("孔候选", "按面三角形拟合圆、法向轴向性、半径误差和对齐度判断内圆柱", "孔数量、孔加工复杂度、2D 复核提示"),
        ("曲面面数", "非平面且未满足圆柱拟合的面", "复杂度和工艺提示"),
        ("薄壁风险", "平均截面 2*volume/area < 最大尺寸 * 0.045", "提示人工复核装夹和变形"),
    ], [1800, 4800, 2760], font_size=8.8)
    add_code(doc, """const metrics = {
  meshCount: result.meshes.length,
  faces, triangles, size: max.map((v, i) => v - min[i]),
  area, volume: Math.abs(volume)
};
metrics.features = analyzeBrepFeatures(result, metrics);

const utilization = Math.max(.01, Math.min(1,
  bounds.volume / bboxVolume));
const thinWallRisk = meanSection > 0
  && meanSection < maxDimension * .045;""", "app.js:52-55（按可读性换行）", language="JavaScript")
    add_note(doc, "工程边界", "这些特征来自三角网格和几何启发式，不是 CAM 刀路验证，也不是制造可行性证明。复杂腔体、非圆柱孔、倒角、公差、夹具和机床能力必须回到释放图纸与工艺评审。", color="9B1C1C", fill="FFF1F1")
    add_heading(doc, "9.3 PDF/FAIR 识别链", 2)
    add_body(doc, "PDF 打开时，pdf.js 同时渲染 canvas 并提取文字框。用户框选区域后，recognize 的顺序是：PDF/vector text 选中 -> 图像符号分析 -> buildVisionCandidates 的字符结构重建 -> OCR fallback -> buildInspectionCandidates -> parseInspectionDimension。设计上的关键约束是“PDF/vector 数字是权威，OCR 只补工程符号和结构证据”，避免 OCR 把 17 误读为 117 后覆盖原始尺寸。")
    add_code(doc, """export function parseInspectionDimension(raw,
  { source = 'PDF 文字层', ocrRaw = '' } = {}) {
  const textRaw = normalizeInspectionText(raw);
  const imageRaw = normalizeOcrEngineeringText(ocrRaw || raw);
  const combined = mergeTextSignals(textRaw, imageRaw);
  const type = detectType(combined);
  const authoritative = imageOnly
    ? (imageRaw || textRaw)
    : preferredTextSignal(textRaw, imageRaw);
  const text = repairEngineeringNotation(
    nominalText(authoritative, type), type, imageRaw);
  const tolerance = toleranceParts(textRaw);
  return { type, text, ...tolerance, recognition: ... };
}""", "inspection-recognition.js:150-160（逻辑摘录）", language="JavaScript")
    add_table(doc, ["识别对象", "规则示例", "人工复核点"], [
        ("直径/孔径", "Ø19、3 x Ø6.6、8.1 H7", "符号是否来自选区，配合图纸基准和公差"),
        ("螺纹", "M6 ↧12；OCR 的 I/T/V 深度字形归一化", "螺纹规格、有效深度、塞规/针规"),
        ("沉孔", "⌴ Ø12 ↧52；上下两行结构恢复", "沉孔与通孔的层级关系"),
        ("粗糙度/形位", "Ra、位置度、平面度、垂直度等类型映射", "基准 A/B/C、测量工具和检验方法"),
        ("普通工程文字", "INSPECTION NOTE、材料说明等保留为 other", "是否应转成独立质量特性"),
    ], [1800, 4400, 3160], font_size=8.8)
    add_heading(doc, "9.4 本地报价模型与导出", 2)
    add_body(doc, "quoteBasis 以材料密度/材料费率、包围盒与实体体积、孔/曲面/薄壁风险、面数和材料利用率估算 machiningHours、machineRate、setupCost，再得到 estimate。partCost 继续叠加材料与毛坯、机加工、装夹/编程、表面处理、首检/质检、制造管理费和利润。用户可以手动锁定单价，导出逻辑会区分自动核算和手动锁定。")
    add_table(doc, ["输出", "内容", "持久化位置"], [
        ("制造 BOM", "零件、材料、工艺、尺寸、体积、重量、利用率、特征和图纸关联", "浏览器下载 XLSX"),
        ("报价交付包", "报价汇总、制造 BOM、成本明细、项目文件、报价说明五个 sheet", "浏览器下载 XLSX"),
        ("正式报价 PDF", "中文字体、含税/运费/条款/签字区", "浏览器下载 PDF"),
        ("云端报价草稿", "quote + quote_lines，金额按 cents 存储", "SQLite；不含 B-Rep 原始指标"),
    ], [1800, 4700, 2860], font_size=8.8)
    add_note(doc, "双模型风险", "本地工作台按几何特征和估算费率计算报价，云端 BOM API 只保存用户输入的数量和单价。两条链没有共享统一的成本规则、版本号或审批状态，因此同一零件可能产生不同报价。正式业务应把成本参数、几何快照、报价版本和审批结果建成可追溯的服务端对象。", color="9A6A00", fill="FFF8E8")

    # 10
    add_heading(doc, "10  ONLYOFFICE 文件链", 1, page_break=True)
    add_heading(doc, "10.1 文档生命周期", 2)
    add_numbered(doc, [
        "POST /api/documents 从 templates/blank.* 复制模板，或 POST /api/documents/upload 接收原始流，生成组织目录下的 storage_key。",
        "GET /api/documents 返回当前用户可见文档，按 owner、scope 和 office_document_permissions 过滤；项目文档还要通过项目访问校验。",
        "用户打开文档时，GET /api/documents/:id/config 生成 document URL、编辑/下载权限、callbackUrl 和 JWT。",
        "onlyoffice.html 加载 DocumentServer API，编辑器保存时调用 /onlyoffice/callback/:documentId；status 2/6 时服务端下载回调 URL 并更新文件大小和 version。",
        "文档删除先进入 deleted_at 回收站，永久删除才从 SQLite 和文件目录移除。",
    ])
    add_heading(doc, "10.2 签名和权限分析", 2)
    add_code(doc, """const signature = fileSignature(session.organizationId, document.id);
const config = {
  documentServer: officeServer,
  document: {
    key: `${document.id}-v${document.version}`,
    title: document.title,
    url: `${publicBase}/office-files/${document.id}/...&signature=${signature}`,
    permissions: { edit: document.permission === 'edit', download: true, print: true }
  },
  editorConfig: { mode: document.permission === 'edit' ? 'edit' : 'view', callbackUrl }
};
config.token = jwt(configWithoutServer);""", "server/app.js:242-255（逻辑摘录）", language="JavaScript")
    add_body(doc, "服务端用 HMAC(secret, organizationId:documentId) 保护文件 URL，再用 JWT 给 ONLYOFFICE 配置签名；文档 key 通过 version 变化避免编辑器缓存旧版本。当前签名没有显式过期时间，URL 一旦泄露，在密钥轮换前可能长期有效；正式环境应加入短期 exp、单次/版本绑定、来源限制和密钥轮换。")
    add_heading(doc, "10.3 关键风险", 2)
    add_table(doc, ["级别", "问题", "修复建议"], [
        ("高", "onlyoffice.js 的本地 fallback 使用硬编码 mfggo-onlyoffice-2026 secret", "生产路径只接受服务端 config；删除浏览器签名 fallback，密钥永不下发客户端"),
        ("高", "callback 直接 fetch body.url 并写文件，缺少大小/超时/来源限制", "限制 DocumentServer 网段、响应大小、Content-Type、超时和幂等键"),
        ("中", "文件按本地目录保存，多实例无法共享", "对象存储 + 版本对象 + 元数据事务"),
        ("中", "PDF 也允许 download/print，细粒度权限未建模", "增加 document.download/document.print 与字段级策略"),
        ("中", "删除/恢复/永久删除操作部分未追加审计", "所有文件状态变更统一写审计并保留操作者"),
    ], [1000, 4200, 4160], font_size=8.7)

    # 11
    add_heading(doc, "11  安全、权限与审计", 1, page_break=True)
    add_heading(doc, "11.1 当前角色矩阵", 2)
    add_table(doc, ["角色", "当前代码能力", "数据范围"], [
        ("owner", "ROLE_PERMISSIONS = ['*']；平台管理员 seed 账号同时为 owner", "全组织、全项目；平台 API 另需 platform_admin"),
        ("admin", "ROLE_PERMISSIONS = ['*']", "全组织、全项目"),
        ("engineer", "项目/零件/报价/任务/沟通/聊天/文档读写；FAIR 读", "必须加入项目；独立任务默认本人"),
        ("qa", "项目/零件/FAIR/任务/沟通/聊天/文档读写", "必须加入项目；无 quote 写权限"),
        ("viewer", "项目/零件/报价/FAIR/任务/沟通/聊天/文档读", "必须加入项目；不可写和不可读审计"),
    ], [1300, 5200, 2860], font_size=8.8)
    add_heading(doc, "11.2 权限判定顺序", 2)
    add_body(doc, "实际请求的推荐解释顺序是：会话有效 -> 角色拥有 permission -> 企业启用 moduleKey -> resource 属于 session.organizationId -> 项目存在且成员可访问 -> 具体字段/动作约束 -> 写入审计。当前 authorize 已覆盖前五步的核心部分，任务负责人、成员角色变更和文件 owner 又补充了专门校验。")
    add_table(doc, ["边界", "已做", "仍需补齐"], [
        ("身份", "scrypt + timingSafeEqual；Bearer session；12h expiration", "MFA、密码策略、登录限流、刷新/撤销、设备管理"),
        ("租户", "organization_id 查询和组织切换会话", "企业状态、跨组织邀请、数据区域、代理会话"),
        ("角色", "五种静态角色；owner/admin 通配", "可配置 RBAC、权限版本、职责分离、报价审批"),
        ("项目", "project_members；聚合接口按权限裁剪", "项目角色动作矩阵、部门/本人范围、外部协作者"),
        ("敏感字段", "quote 访问受 role/module/project 约束", "成本、毛利、客户联系方式、原始图纸下载的字段级权限"),
        ("审计", "业务写操作与登录等动作记录 audit_events", "不可篡改存储、分页归档、IP/设备、前后值和检索"),
    ], [1600, 3800, 3960], font_size=8.7)
    add_heading(doc, "11.3 需要优先修复的代码级问题", 2)
    add_table(doc, ["优先级", "发现", "建议动作"], [
        ("P0", "README 和 env 示例仍出现 admin/123456；seed 有默认密码回退", "首次启动强制随机密码或一次性初始化；启动后禁止弱口令"),
        ("P0", "platform 创建企业时把当前 platform admin 直接作为 owner membership", "平台角色与企业 owner 分离；使用邀请/受控代理会话"),
        ("P0", "onlyoffice.js 浏览器 fallback 硬编码 JWT secret", "删除客户端 secret 和 demo fallback，生产必须由后端签名"),
        ("P1", "quote/fair 的 partId 未验证同组织、同项目", "增加资源归属检查和跨租户测试"),
        ("P1", "报价头与明细写入没有事务", "封装 transaction(fn)，失败回滚并增加故障测试"),
        ("P1", "文件回调和上传缺少全局限流、超时、病毒扫描", "接入限流、内容识别、扫描队列和对象存储"),
    ], [1000, 4900, 3460], font_size=8.7)

    # 12
    add_heading(doc, "12  测试与质量", 1, page_break=True)
    add_heading(doc, "12.1 本次实际验证结果", 2)
    add_note(doc, "测试结果", "使用 Node.js v24.19.0 运行 `node --test tests/saas-api.test.mjs tests/zhizao-app.test.mjs tests/inspection-recognition.test.mjs`，共 82 个测试，82 通过，0 失败，0 跳过；总耗时约 8.2 秒。", color="0F6B78", fill="EAF6F3")
    add_body(doc, "构建验证：`pnpm run build` 成功，Vite 7.1.1 转换 304 个模块；同时报告 workspace 产物约 1,997.55 kB、gzip 580.94 kB 的大 chunk 警告，后续应通过动态导入或 manualChunks 拆分 CAD、PDF/OCR 和导出能力。", after=6)
    add_table(doc, ["测试文件", "数量/范围", "验证内容"], [
        ("tests/saas-api.test.mjs", "25 个", "临时 SQLite、登录、租户隔离、角色拒绝、模块关闭、项目成员、任务、文档、报价、FAIR、聊天附件、审计、workspace 聚合"),
        ("tests/zhizao-app.test.mjs", "21 个", "store、导航、筛选、平台切换、移动端控制、文档入口、聊天状态、旧租户清理、异步竞态"),
        ("tests/inspection-recognition.test.mjs", "36 个", "尺寸、公差、工程符号、OCR 证据、聚类、结构恢复、噪声抑制、历史数据迁移"),
    ], [2300, 1800, 5260], font_size=8.8)
    add_heading(doc, "12.2 测试设计的强项", 2)
    add_bullets(doc, [
        "SaaS API 使用临时数据库和临时附件目录，测试不会污染真实 data/machquote.sqlite。",
        "权限测试不仅验证 403，还验证拒绝原因和被加入项目后的授权变化。",
        "模块关闭测试验证 API 层拒绝，而不是只验证前端导航消失。",
        "OCR 测试明确验证 PDF 数字不可被错误 OCR 覆盖，这是工程图识别最关键的安全约束之一。",
        "前端测试用源码断言锁住了租户清理、旧路由删除和异步选中项保护等回归点。",
    ])
    add_heading(doc, "12.3 尚未覆盖的质量面", 2)
    add_table(doc, ["缺口", "为什么重要", "建议测试"], [
        ("浏览器端到端", "Three/WebGL、WASM、pdf.js canvas、真实 OCR 在 Node 测试中未执行", "Playwright + 样例 STEP/PDF + 截图像素/导出文件校验"),
        ("故障恢复", "进程崩溃、SQLite 锁、磁盘满、callback 重试未覆盖", "故障注入、重启恢复、幂等和备份恢复演练"),
        ("性能", "大模型、长会话、1000+ 项审计和并发上传未基准化", "基准数据集、内存峰值、P95 延迟和并发写压测"),
        ("安全扫描", "依赖、文件内容、CSRF、限流、SSRF 没有自动门禁", "CI 依赖审计、DAST、恶意文件和 URL 测试"),
    ], [1800, 3900, 3660], font_size=8.8)
    add_heading(doc, "12.4 建议的 CI 门禁", 2)
    add_numbered(doc, [
        "pnpm install --frozen-lockfile 与 pnpm run build。",
        "Node 内置三组测试，要求 0 fail。",
        "启动临时服务，执行 health、登录、租户隔离和文件回调 smoke test。",
        "对文档、依赖和上传路径执行安全扫描；禁止示例弱密码和硬编码 secret 进入生产配置。",
        "对关键样例 CAD/PDF 运行浏览器回归，检查导出 XLSX/PDF 可打开且字段不为空。",
    ])

    # 13
    add_heading(doc, "13  风险、技术债与改进路线", 1, page_break=True)
    add_heading(doc, "13.1 分阶段路线", 2)
    roadmap = [
        ("P0 立即", "凭证与文件安全", "移除默认/客户端 secret；强制初始化改密；HTTPS/VPN；上传限流、大小、内容识别；callback 超时/来源/幂等"),
        ("P1", "企业与成员生命周期", "organization status；invitation token；首次登录改密；成员停用/移除；密码重置；MFA；session revoke"),
        ("P2", "可配置数据权限", "roles/permissions/role_permissions；项目角色；部门/本人范围；quote.read_cost；审批动作；外部协作者"),
        ("P3", "制造业务闭环", "客户/供应商、询盘、订单、工艺路线、排产、交付、版本和归档；报价与 FAIR 版本化"),
        ("P4", "规模化部署", "PostgreSQL、对象存储、队列、备份恢复、监控告警、数据保留、SSO、审计归档"),
    ]
    add_table(doc, ["阶段", "主题", "交付"], roadmap, [1200, 2300, 5860], font_size=8.8)
    add_heading(doc, "13.2 推荐的重构顺序", 2)
    add_body(doc, "第一步不是重写前端，而是把认证、租户上下文、授权和资源归属检查抽成可测试的 policy/domain 层；第二步把 quote/fair/document 的版本、事务和审计补齐；第三步再拆分 Koa 路由和前端 store。这样可以先稳住安全和数据正确性，再降低单文件维护成本。")
    add_table(doc, ["重构目标", "当前做法", "目标做法"], [
        ("授权", "ROLE_PERMISSIONS 常量 + authorize", "Policy service + DB 配置 RBAC + resource scope"),
        ("数据库访问", "db.js 一个对象承载所有域方法", "按域拆 repository/service；统一 tenantScope()"),
        ("报价", "浏览器估算 + 云端草稿两套模型", "服务端成本参数、geometry snapshot、quote_version、approval"),
        ("文件", "本地目录 storage_key", "对象存储版本对象 + 元数据事务 + signed URL exp"),
        ("前端", "单文件 render/事件委托", "按域拆 store、纯函数、异步资源缓存和页面组件"),
        ("可观测性", "console.error + audit_events", "结构化日志、request ID、指标、告警和审计归档"),
    ], [1800, 3900, 3660], font_size=8.7)
    add_heading(doc, "13.3 验收优先级建议", 2)
    add_body(doc, "如果目标是继续内部验收，优先修复默认凭证、ONLYOFFICE 客户端 secret、报价/FAIR 关联归属和报价事务，然后补一套 Playwright 样例。若目标是正式公网客户，必须在此基础上完成 HTTPS、MFA/限流、对象存储、备份恢复和 PostgreSQL 迁移评估；不能因为 82/82 单测通过就把当前版本定义为生产 SaaS。")
    add_note(doc, "最终判断", "当前版本适合“可运行的机加工报价协同基础”和“本地 CAD/FAIR 体验验证”，不适合未经整改直接承载正式客户的高敏感图纸、成本和报价审批。", color="9B1C1C", fill="FFF1F1")

    # 14
    add_heading(doc, "14  附录", 1, page_break=True)
    add_heading(doc, "14.1 关键文件索引", 2)
    inventory = [
        ("server/index.js", source_line_count("server/index.js"), "服务启动、环境校验、静态目录和关闭信号"),
        ("server/app.js", source_line_count("server/app.js"), "Koa、认证、授权、API、文件和 ONLYOFFICE 回调"),
        ("server/db.js", source_line_count("server/db.js"), "SQLite schema、seed、迁移、查询映射和审计"),
        ("zhizao-app.js", source_line_count("zhizao-app.js"), "云端工作台 store、渲染、导航、事件和 API 调用"),
        ("app.js", source_line_count("app.js"), "本地 CAD、PDF/FAIR、报价计算和导出"),
        ("inspection-recognition.js", source_line_count("inspection-recognition.js"), "工程尺寸类型、OCR 修复、解析、聚类和历史迁移"),
        ("inspection-pipeline.js", source_line_count("inspection-pipeline.js"), "vector/OCR/视觉候选组装和工程结构语法"),
        ("workspace.html", source_line_count("workspace.html"), "本地 3D/FAIR 页面骨架和依赖入口"),
        ("vite.config.js", source_line_count("vite.config.js"), "多页面构建和 API 代理"),
        ("ARCHITECTURE.md", source_line_count("ARCHITECTURE.md"), "目标企业体系、权限模型和路线图"),
        ("README.md", source_line_count("README.md"), "启动、ONLYOFFICE、验收、公网部署说明"),
    ]
    add_table(doc, ["文件", "行数", "职责"], [(a, str(b), c) for a, b, c in inventory], [2600, 900, 5860], font_size=8.8)
    add_heading(doc, "14.2 常用验收命令", 2)
    add_code(doc, """pnpm install
pnpm run build
pnpm run server

# 开发联调
pnpm run dev

# 回归测试
node --test tests/saas-api.test.mjs \
  tests/zhizao-app.test.mjs \
  tests/inspection-recognition.test.mjs

# 健康检查
curl http://127.0.0.1:4310/api/health""", "README.md:7-32; 本次验证命令", language="PowerShell / Bash")
    add_heading(doc, "14.3 接口快速索引", 2)
    quick_api = [
        ("认证", "POST /api/auth/login", "用户名/密码 -> token"),
        ("当前上下文", "GET /api/me", "用户、组织、角色、模块、平台标志"),
        ("项目", "GET/POST /api/projects; PUT /api/projects/:projectId", "项目看板和项目更新"),
        ("聚合", "GET /api/projects/:projectId/workspace", "按权限裁剪的项目资源"),
        ("零件", "GET /api/parts; POST /api/projects/:projectId/parts; PUT /api/parts/:partId", "零件元数据"),
        ("报价", "GET/POST/PUT /api/projects/:projectId/quotes", "报价头和 BOM 明细"),
        ("FAIR", "GET/POST/PUT /api/projects/:projectId/fair-items", "检验特性与状态"),
        ("任务", "GET/POST/PUT /api/tasks; /detail; /subtasks; /comments", "任务协同"),
        ("沟通", "GET/POST /api/conversations; /messages; PUT /state", "项目/企业聊天"),
        ("文档", "GET/POST /api/documents; /upload; /config; /download", "Office 文档中心"),
        ("审计", "GET /api/audit; GET /api/platform/audit", "企业/平台审计"),
    ]
    add_table(doc, ["域", "路径", "用途"], quick_api, [1500, 5100, 2760], font_size=8.8)
    add_heading(doc, "14.4 术语表", 2)
    glossary = [
        ("organization_id", "企业租户标识；业务查询和写入的第一隔离条件。"),
        ("project_members", "企业成员参与具体项目的授权关系；非 owner/admin 访问项目的必要条件。"),
        ("FAIR", "First Article Inspection Report，首件检验/检验特性记录。"),
        ("B-Rep", "Boundary Representation，CAD 实体边界表示；本项目通过 OpenCascade WASM 读取并网格化。"),
        ("OCR evidence", "图像识别提供的符号/文字证据；在当前规则中不能覆盖完整 PDF 数字。"),
        ("workspace aggregation", "项目工作台聚合接口；服务端按模块和权限裁剪子资源。"),
        ("P0/P1/P2/P3/P4", "从权限收口、成员生命周期、数据权限、制造闭环到企业级部署的路线阶段。"),
    ]
    add_table(doc, ["术语", "定义"], glossary, [2200, 7160], font_size=9)
    add_note(doc, "文档结束", "本报告只描述当前工作区可验证的代码和明确标注的改进建议。任何上线决策都应以实际部署环境、客户数据分类和安全验收结果为准。", color="0F6B78")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
